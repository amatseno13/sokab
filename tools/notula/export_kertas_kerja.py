#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
=============================================================
EXPORT KERTAS KERJA (SOKAB → Excel)
=============================================================
Mengisi template kertas_kerja.xlsx dengan data Entry Capaian Kinerja.

PRINSIP: hanya menulis ke sel yang TIDAK berisi formula.
Setiap sel diperiksa dulu; kalau isinya diawali '=', sel dilewati dan
dicatat di daftar 'dilewati'. Jadi seluruh rumus di kertas kerja tetap utuh
dan akan menghitung ulang sendiri saat file dibuka.

Pemakaian:
    python3 export_kertas_kerja.py \
        --json /tmp/data.json \
        --template template/kertas_kerja.xlsx \
        --output /tmp/hasil.xlsx

Struktur JSON yang diharapkan:
{
  "satker": "...", "nilai_sakip": "74.45", "tahun": 2026,
  "iku_list": [
     {"kode":"1.1.1.1", "jenis_satuan":"%",
      "x_tw1":..,"x_tw2":..,"x_tw3":..,"x_tw4":..,
      "y_tw1":..,"y_tw2":..,"y_tw3":..,"y_tw4":..,
      "realisasi_tw1":..,  (untuk Non %)
      "kendala":"","solusi":"","rtl":"","pic":"","batas":"",
      "link_bukti":"","link_tl":""}
  ]
}

Mencetak satu baris JSON ke stdout.
=============================================================
"""

import os
import sys
import json
import argparse

try:
    from openpyxl import load_workbook
except ImportError:
    sys.stderr.write("ERROR: openpyxl belum terinstall. Jalankan: pip3 install openpyxl\n")
    sys.exit(2)


SHEET = "LK_Kabkot"

# Peta kolom di sheet LK_Kabkot (1-based)
COL_KODE      = 4    # D
COL_JENIS     = 8    # H  ("IKU" / "Proksi")
COL_SATUAN    = 10   # J  ("%" / "Non %")
COL_TARGET    = 11   # K
COL_ALOKASI   = [13, 14, 15, 16]   # M–P  target kumulatif TW I–IV
COL_REALISASI = [17, 18, 19, 20]   # Q–T  realisasi kumulatif TW I–IV
COL_KENDALA   = 29   # AC
COL_SOLUSI    = 30   # AD
COL_RTL       = 31   # AE
COL_PIC       = 32   # AF
COL_BATAS     = 33   # AG
COL_LINK      = 34   # AH
COL_LINK_TL   = 35   # AI

# Sel header
SEL_SATKER = (3, 5)   # E3
SEL_SAKIP  = (4, 5)   # E4  (E5 predikat adalah FORMULA — jangan disentuh)


class Penulis:
    """Menulis sel hanya bila bukan formula. Mencatat semua yang dilewati."""

    def __init__(self, ws):
        self.ws = ws
        self.ditulis = 0
        self.dilewati = []

    @staticmethod
    def _formula(v):
        return isinstance(v, str) and v.startswith("=")

    def tulis(self, row, col, nilai, label=""):
        if nilai is None or nilai == "":
            return False
        sel = self.ws.cell(row, col)
        if self._formula(sel.value):
            self.dilewati.append({
                "sel": sel.coordinate,
                "label": label,
                "alasan": "berisi formula",
            })
            return False
        sel.value = nilai
        self.ditulis += 1
        return True


def angka(v):
    """Ubah ke float bila memang angka; kalau bukan, kembalikan apa adanya."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", ".")
    try:
        f = float(s)
        return int(f) if f == int(f) else f
    except ValueError:
        return v


def cari_baris_iku(ws):
    """Petakan kode IKU → nomor baris di sheet."""
    peta = {}
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, COL_JENIS).value or "").strip().upper() == "IKU":
            kode = str(ws.cell(r, COL_KODE).value or "").strip()
            if kode:
                peta[kode] = r
    return peta


def punya_xy(ws, baris):
    """
    IKU berjenis '%' punya dua baris turunan: X (baris+1) dan Y (baris+2).
    Dideteksi dari label di kolom E, bukan diasumsikan.
    """
    lx = str(ws.cell(baris + 1, 5).value or "").strip().upper()
    ly = str(ws.cell(baris + 2, 5).value or "").strip().upper()
    return lx.startswith("X") and ly.startswith("Y")


def isi(data, template, output):
    wb = load_workbook(template)
    if SHEET not in wb.sheetnames:
        raise ValueError("Sheet '%s' tidak ada di template. Sheet tersedia: %s"
                         % (SHEET, wb.sheetnames))
    ws = wb[SHEET]
    p = Penulis(ws)

    # ── Header ──
    p.tulis(*SEL_SATKER, data.get("satker"), "Satuan Kerja")
    p.tulis(*SEL_SAKIP, angka(data.get("nilai_sakip")), "Nilai SAKIP")
    # E5 (predikat) sengaja tidak disentuh: isinya formula turunan dari E4.

    peta = cari_baris_iku(ws)
    terisi, tak_ketemu = [], []

    for iku in data.get("iku_list", []):
        kode = str(iku.get("kode", "")).strip()
        baris = peta.get(kode)
        if baris is None:
            tak_ketemu.append(kode)
            continue

        # ── Angka ──
        if str(iku.get("jenis_satuan", "%")).strip() == "%" and punya_xy(ws, baris):
            # Realisasi disimpan sebagai pembilang (X) dan penyebut (Y)
            bx, by = baris + 1, baris + 2
            for i in range(4):
                p.tulis(bx, COL_REALISASI[i], angka(iku.get("x_tw%d" % (i + 1))),
                        "%s X realisasi TW%d" % (kode, i + 1))
                p.tulis(by, COL_REALISASI[i], angka(iku.get("y_tw%d" % (i + 1))),
                        "%s Y realisasi TW%d" % (kode, i + 1))
        else:
            # Non % — realisasi langsung di baris IKU
            for i in range(4):
                p.tulis(baris, COL_REALISASI[i], angka(iku.get("realisasi_tw%d" % (i + 1))),
                        "%s realisasi TW%d" % (kode, i + 1))

        # ── Narasi & tautan (selalu di baris IKU) ──
        for col, kunci, label in [
            (COL_KENDALA, "kendala",    "Kendala"),
            (COL_SOLUSI,  "solusi",     "Solusi"),
            (COL_RTL,     "rtl",        "RTL"),
            (COL_PIC,     "pic",        "PIC"),
            (COL_BATAS,   "batas",      "Batas waktu"),
            (COL_LINK,    "link_bukti", "Link bukti"),
            (COL_LINK_TL, "link_tl",    "Link TL"),
        ]:
            p.tulis(baris, col, iku.get(kunci), "%s %s" % (kode, label))

        terisi.append(kode)

    # Paksa Excel menghitung ulang seluruh formula saat file dibuka,
    # karena nilai hasil hitungan lama tidak ikut tersimpan.
    wb.calculation.fullCalcOnLoad = True
    wb.save(output)

    return {
        "sel_ditulis": p.ditulis,
        "sel_dilewati": len(p.dilewati),
        "dilewati": p.dilewati[:20],
        "iku_terisi": terisi,
        "iku_tak_ketemu": tak_ketemu,
    }


def main():
    ap = argparse.ArgumentParser(description="Export kertas kerja SOKAB ke Excel")
    ap.add_argument("--json",     required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--output",   required=True)
    args = ap.parse_args()

    for f in (args.json, args.template):
        if not os.path.exists(f):
            sys.stderr.write("ERROR: file tidak ditemukan: %s\n" % f)
            sys.exit(2)

    with open(args.json, "r", encoding="utf-8") as fh:
        data = json.load(fh)

    try:
        info = isi(data, args.template, args.output)
    except Exception as e:
        sys.stderr.write("ERROR saat mengisi Excel: %s\n" % e)
        sys.exit(3)

    info.update({"success": True, "output": args.output})
    print(json.dumps(info, ensure_ascii=False))


if __name__ == "__main__":
    main()
