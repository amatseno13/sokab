#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
=============================================================
GENERATE NOTULA MONITORING KINERJA TRIWULANAN  (SOKAB Edition)
=============================================================
Dipanggil dari aplikasi SOKAB, bukan lagi diedit manual.

MODE 1 — dari Excel yang di-upload user (dipakai SOKAB):
    python3 generate_notula.py \
        --excel /tmp/upload.xlsx --sheet LK_Kabkot --triwulan II \
        --meta /tmp/meta.json \
        --template template/word_FRA.docx \
        --output /tmp/hasil.docx

MODE 2 — hanya baca Excel, tanpa membuat dokumen (untuk preview):
    python3 generate_notula.py --excel /tmp/upload.xlsx --triwulan II --inspect

MODE 3 — dari JSON (misal data langsung dari database SOKAB):
    python3 generate_notula.py --json data.json --template t.docx --output out.docx

Semua mode mencetak SATU baris JSON ke stdout.
Error dicetak ke stderr dengan exit code != 0.
=============================================================
"""

import re
import os
import sys
import json
import argparse

try:
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_LINE_SPACING
except ImportError:
    sys.stderr.write("ERROR: python-docx belum terinstall. "
                     "Jalankan: pip3 install python-docx openpyxl\n")
    sys.exit(2)


TW_LIST = ['I', 'II', 'III', 'IV']

FIELD_IKU = ("kode", "nama", "sasaran_kode", "sasaran_nama", "tujuan",
             "target_pk", "satuan", "alokasi_tw", "real_tw", "capaian_tw",
             "capaian_pk", "kendala", "solusi", "rtl", "pic", "batas",
             "link_bukti", "link_tl_sblm", "ro_narasi")


def normalisasi(data):
    """Pastikan semua field ada dan bertipe string."""
    data.setdefault("satker", "BPS Kabupaten/Kota")
    data.setdefault("nilai_sakip", "")
    data.setdefault("predikat", "")
    data.setdefault("triwulan", "I")
    data.setdefault("tahun", 2026)
    data.setdefault("agenda", {})
    data.setdefault("iku_list", [])
    for iku in data["iku_list"]:
        for k in FIELD_IKU:
            v = iku.get(k)
            iku[k] = "" if v is None else str(v)
    return data


# ============================================================
# SUMBER DATA 1 — JSON
# ============================================================
def baca_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return normalisasi(json.load(f))


# ============================================================
# SUMBER DATA 2 — EXCEL (file yang di-upload user)
# ============================================================
def pilih_sheet(wb, diminta):
    """Cari sheet yang diminta; kalau tidak ada, tebak yang paling masuk akal."""
    if diminta and diminta in wb.sheetnames:
        return diminta
    for nama in wb.sheetnames:                      # cocokkan tanpa peduli huruf besar/kecil
        if diminta and nama.strip().lower() == diminta.strip().lower():
            return nama
    for nama in wb.sheetnames:                      # sheet kertas kerja biasanya diawali "LK"
        if nama.strip().upper().startswith("LK"):
            return nama
    return wb.sheetnames[0]


def baca_excel(path, sheet, triwulan=None):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.stderr.write("ERROR: openpyxl belum terinstall.\n")
        sys.exit(2)

    try:
        wb = load_workbook(path, data_only=True, read_only=False)
    except Exception as e:
        sys.stderr.write("ERROR: file Excel tidak bisa dibaca (%s). "
                         "Pastikan formatnya .xlsx, bukan .xls lama.\n" % e)
        sys.exit(2)

    sheet_dipakai = pilih_sheet(wb, sheet)
    ws = wb[sheet_dipakai]

    # --- Info header (kolom E) ---
    satker      = ws.cell(3, 5).value or "BPS Kabupaten/Kota"
    nilai_sakip = ws.cell(4, 5).value
    predikat    = ws.cell(5, 5).value

    # --- Tentukan triwulan ---
    if triwulan and str(triwulan).upper() in TW_LIST:
        tw = str(triwulan).upper()
    else:
        tw_raw = ws.cell(6, 29).value or ""
        m = re.search(r'\b(IV|III|II|I)\b', str(tw_raw).upper())
        tw = m.group(1) if m else "I"
    tw_idx = TW_LIST.index(tw)          # 0..3

    # --- Offset kolom per triwulan (blok 4 kolom TW I..IV) ---
    # 0-based index pada tuple row:
    #   alokasi target kumulatif : M..P  -> 12..15
    #   realisasi kumulatif      : Q..T  -> 16..19
    #   capaian thd target TW    : U..X  -> 20..23
    #   capaian thd target PK    : Y..AB -> 24..27
    i_alokasi = 12 + tw_idx
    i_real    = 16 + tw_idx
    i_cap_tw  = 20 + tw_idx
    i_cap_pk  = 24 + tw_idx

    def val(v):
        if v is None:
            return ""
        if isinstance(v, float):
            if v == int(v):
                return str(int(v))
            return ("%.2f" % v).rstrip('0').rstrip('.')
        return str(v).strip()

    def cell(row, i):
        return val(row[i]) if len(row) > i else ""

    iku_list, tujuan, sasaran_kode, sasaran_nama = [], "", "", ""
    peringatan = []

    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 10 or i > 400:
            continue
        r0 = str(row[0]) if row[0] else ""

        # Baris Tujuan
        if row[0] and re.match(r'^T\d\s*:', r0.strip()):
            tujuan = r0.strip()

        # Baris Sasaran (kode di kolom B, nama di kolom C)
        elif len(row) > 7 and row[1] and row[2] and row[7] is None:
            sasaran_kode = str(row[1]).strip()
            sasaran_nama = str(row[2]).strip()

        # Baris IKU (kode di kolom D, kolom H = "IKU")
        elif len(row) > 7 and row[3] and str(row[7]).strip().upper() == "IKU":
            iku_list.append({
                "kode": str(row[3]).strip(),
                "nama": str(row[4]).strip() if row[4] else "",
                "sasaran_kode": sasaran_kode,
                "sasaran_nama": sasaran_nama,
                "tujuan": tujuan,
                "target_pk":  cell(row, 10),
                "satuan":     cell(row, 11),
                "alokasi_tw": cell(row, i_alokasi),
                "real_tw":    cell(row, i_real),
                "capaian_tw": cell(row, i_cap_tw),
                "capaian_pk": cell(row, i_cap_pk),
                "kendala":      cell(row, 28),
                "solusi":       cell(row, 29),
                "rtl":          cell(row, 30),
                "pic":          cell(row, 31),
                "batas":        cell(row, 32),
                "link_bukti":   cell(row, 33),
                "link_tl_sblm": cell(row, 34),
                "ro_narasi":    cell(row, 39),
            })

    if not iku_list:
        peringatan.append(
            "Tidak ada baris IKU yang terbaca dari sheet '%s'. "
            "Pastikan kolom D berisi kode IKU dan kolom H berisi teks 'IKU'." % sheet_dipakai)

    return normalisasi({
        "satker": str(satker).strip(),
        "nilai_sakip": val(nilai_sakip),
        "predikat": val(predikat),
        "triwulan": tw,
        "tahun": 2026,
        "agenda": {},
        "iku_list": iku_list,
        "_sheet": sheet_dipakai,
        "_sheets_tersedia": wb.sheetnames,
        "_peringatan": peringatan,
    })


# ============================================================
# GABUNG METADATA (form di SOKAB menimpa nilai dari Excel)
# ============================================================
def gabung_meta(data, meta_path):
    with open(meta_path, "r", encoding="utf-8") as f:
        meta = json.load(f)

    for k in ("satker", "nilai_sakip", "predikat", "triwulan"):
        if str(meta.get(k, "")).strip():
            data[k] = str(meta[k]).strip()
    if meta.get("tahun"):
        data["tahun"] = meta["tahun"]

    agenda = data.get("agenda") or {}
    agenda.update({k: v for k, v in (meta.get("agenda") or {}).items() if str(v or "").strip()})
    data["agenda"] = agenda
    return data


# ============================================================
# UTILITAS DOCX
# ============================================================
def unique_cells(row):
    """Kembalikan sel unik saja (hindari duplikat akibat merged cell)."""
    seen, result = set(), []
    for cell in row.cells:
        cid = id(cell._tc)
        if cid not in seen:
            seen.add(cid)
            result.append(cell)
    return result


def bersihkan_nilai(text):
    """Ambil baris pertama & buang pengulangan string ganda."""
    if not text:
        return ""
    text = str(text).strip()
    baris = [b.strip() for b in text.split("\n") if b.strip()]
    if not baris:
        return text
    hasil = baris[0]
    words = hasil.split(" ")
    for length in range(1, len(words) // 2 + 1):
        chunk = " ".join(words[:length])
        rest = hasil[len(chunk):].strip()
        if rest.startswith(chunk):
            return chunk
    return hasil


def set_cell_text(cell, text, bold=False, font_size=None):
    for para in cell.paragraphs:
        for run in para.runs:
            run.text = ""
    para = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    run = para.add_run(str(text) if text else "")
    run.bold = bold
    run.font.name = "Arial"
    if font_size:
        run.font.size = Pt(font_size)
    return run


def buat_para_nilai(nilai, is_link=False):
    """Paragraf nilai: Arial, tidak bold. Link = rata kiri, teks biasa = justify."""
    new_p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    jc = OxmlElement("w:jc")
    if is_link:
        jc.set(qn("w:val"), "left")
        teks = str(nilai).rstrip() + " "
    else:
        jc.set(qn("w:val"), "both")
        teks = str(nilai)
    pPr.append(jc)
    new_p.append(pPr)

    r_el = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")
    rPr.append(rFonts)
    r_el.append(rPr)
    t_el = OxmlElement("w:t")
    t_el.set(qn("xml:space"), "preserve")
    t_el.text = teks
    r_el.append(t_el)
    new_p.append(r_el)
    return new_p


def sisip_nilai_setelah_label(cell, label, nilai, is_link=False):
    """Sisipkan paragraf nilai tepat setelah paragraf berlabel (bold label tetap aman)."""
    if not nilai:
        return False
    for para in cell.paragraphs:
        if label.lower() in para.text.lower():
            para._p.addnext(buat_para_nilai(nilai, is_link=is_link))
            return True
    return False


def replace_text_in_para(para, old, new):
    if old not in para.text:
        return
    for run in para.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
    if old in para.text:                       # teks terpecah antar run
        new_text = para.text.replace(old, new)
        for run in para.runs:
            run.text = ""
        if para.runs:
            para.runs[0].text = new_text
        else:
            para.add_run(new_text)


def replace_in_doc(doc, replacements):
    def walk_cell(cell):
        for para in cell.paragraphs:
            for old, new in replacements.items():
                replace_text_in_para(para, old, str(new))
        for nested in cell.tables:
            for nrow in nested.rows:
                for ncell in nrow.cells:
                    walk_cell(ncell)

    for para in doc.paragraphs:
        for old, new in replacements.items():
            replace_text_in_para(para, old, str(new))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                walk_cell(cell)


def get_all_text_in_table(table):
    return " ".join(cell.text for row in table.rows for cell in row.cells)


def set_arial_seluruh_doc(doc):
    """Arial + line spacing 1.15 + space after 6pt untuk seluruh dokumen."""
    def set_para_format(para):
        for run in para.runs:
            run.font.name = "Arial"
            rPr = run._r.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.insert(0, rFonts)
            rFonts.set(qn("w:ascii"), "Arial")
            rFonts.set(qn("w:hAnsi"), "Arial")
            rFonts.set(qn("w:cs"), "Arial")
        pf = para.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.15
        pf.space_after = Pt(6)

    for para in doc.paragraphs:
        set_para_format(para)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    set_para_format(para)
                for nested in cell.tables:
                    for nrow in nested.rows:
                        for ncell in nrow.cells:
                            for para in ncell.paragraphs:
                                set_para_format(para)


# ============================================================
# AGENDA & TANDA TANGAN
# ============================================================
def isi_agenda(doc, agenda):
    peta = {
        "hari/tanggal":   agenda.get("hari_tanggal", ""),
        "waktu":          agenda.get("waktu", ""),
        "tempat":         agenda.get("tempat", ""),
        "pimpinan rapat": agenda.get("pimpinan", ""),
    }
    if not doc.tables:
        return
    for row in doc.tables[0].rows:
        cells = unique_cells(row)
        if len(cells) < 2:
            continue
        label = cells[0].text.strip().lower()
        for kunci, nilai in peta.items():
            if label == kunci and nilai:
                set_cell_text(cells[1], nilai)


def isi_tanda_tangan(doc, agenda):
    if not doc.tables:
        return
    kepala     = agenda.get("kepala_satker", "")
    notulis    = agenda.get("notulis", "")
    tempat_tgl = agenda.get("tempat_tanggal_ttd", "")

    for row in doc.tables[-1].rows:
        for cell in unique_cells(row):
            teks = cell.text
            if "Kepala Satker" in teks and kepala:
                for para in cell.paragraphs:
                    if para.text.strip() == "xxx":
                        for run in para.runs:
                            run.text = ""
                        (para.runs[0] if para.runs else para.add_run()).text = kepala
                        break
            if "Notulis" in teks:
                if tempat_tgl:
                    for para in cell.paragraphs:
                        if "Tempat, Tanggal" in para.text:
                            replace_text_in_para(para, "Tempat, Tanggal", tempat_tgl)
                            break
                if notulis:
                    for para in cell.paragraphs:
                        if para.text.strip() == "xxx":
                            for run in para.runs:
                                run.text = ""
                            (para.runs[0] if para.runs else para.add_run()).text = notulis
                            break


# ============================================================
# GENERATE DOKUMEN
# ============================================================
def kode_ke_id(kode):
    """'1.1.1.1' -> '1111' (penanda tabel di template)."""
    return str(kode).replace(".", "")


def periksa_template(doc):
    """
    Pastikan file template masih berupa dokumen kosong, bukan notula yang sudah jadi.
    Kalau notula hasil generate dipakai ulang sebagai template, hasilnya kacau:
    nilai lama tidak tertimpa dan label terisi dua kali.
    """
    teks_awal = " ".join(p.text for p in doc.paragraphs[:8])
    if "TRIWULAN XX" not in teks_awal.upper():
        return ("File template sepertinya bukan dokumen kosong — penanda "
                "'TRIWULAN XX' tidak ditemukan, judulnya sudah terisi triwulan tertentu. "
                "Sepertinya yang dipakai adalah notula hasil generate, bukan template asli. "
                "Ganti tools/notula/template/word_FRA.docx dengan file template resmi yang masih kosong.")
    return None


def generate_notula(data, template_path, output_path, abaikan_cek=False):
    doc = Document(template_path)

    if not abaikan_cek:
        masalah = periksa_template(doc)
        if masalah:
            raise ValueError(masalah)

    satker   = data["satker"]
    tw       = data["triwulan"]
    tahun    = str(data.get("tahun", 2026))
    iku_list = data["iku_list"]
    agenda   = data.get("agenda") or {}

    tw_angka = {"I": "pertama", "II": "kedua", "III": "ketiga", "IV": "keempat"}

    # ---- 1. PLACEHOLDER GLOBAL ----
    replacements = {
        "MONITORING KINERJA TRIWULAN XX TAHUN 2026":
            "MONITORING KINERJA TRIWULAN %s TAHUN %s" % (tw, tahun),
        "BPS PROVINSI xxxxxx/ BPS KABUPATEN/KOTA yyy":  satker.upper(),
        "BPS PROVINSI xxxxxx / BPS KABUPATEN/KOTA yyy": satker.upper(),
        "Monitoring Kinerja Triwulan XX Tahun 2026":
            "Monitoring Kinerja Triwulan %s Tahun %s" % (tw, tahun),
        "BPS Provinsi xxx / BPS Kabupaten/kota yyy": satker,
        "Capaian Kinerja IKU triwulan xx tahun 2026 pada BPS Provinsi xxx / "
        "BPS Kabupaten/kota yyy sebesar … persen.":
            "Capaian Kinerja IKU triwulan %s tahun %s pada %s."
            % (tw_angka.get(tw, tw), tahun, satker),
        "Capaian Kinerja IKU triwulan xx tahun 2026 pada BPS Provinsi xxx / "
        "BPS Kabupaten/kota yyy sebesar ... persen.":
            "Capaian Kinerja IKU triwulan %s tahun %s pada %s."
            % (tw_angka.get(tw, tw), tahun, satker),
        "Capaian Kinerja Triwulan XX Tahun 2026":
            "Capaian Kinerja Triwulan %s Tahun %s" % (tw, tahun),
        "Triwulan XX Tahun 2026": "Triwulan %s Tahun %s" % (tw, tahun),
        "Triwulan XX":            "Triwulan %s" % tw,
        "triwulan xx tahun 2026": "triwulan %s tahun %s" % (tw_angka.get(tw, tw), tahun),
    }
    replace_in_doc(doc, replacements)

    # ---- 2. AGENDA & TANDA TANGAN ----
    isi_agenda(doc, agenda)
    isi_tanda_tangan(doc, agenda)

    # ---- 3. ISI DATA PER IKU ----
    tabel_terpakai = set()
    terisi = 0
    tidak_ketemu = []

    for iku in iku_list:
        id_tabel  = kode_ke_id(iku["kode"])
        kode_asli = str(iku["kode"])

        # Template resmi memakai penanda tanpa titik ("1111"), tetapi sebagian
        # salinan sudah memakai kode bertitik — terima keduanya.
        target_table = None
        penanda = id_tabel
        for bentuk in (id_tabel, kode_asli):
            for idx, table in enumerate(doc.tables):
                if idx in tabel_terpakai:
                    continue
                if bentuk in get_all_text_in_table(table):
                    target_table = table
                    penanda = bentuk
                    tabel_terpakai.add(idx)
                    break
            if target_table is not None:
                break

        if target_table is None:
            tidak_ketemu.append(kode_asli)
            continue
        terisi += 1

        for row in target_table.rows:
            ucells = unique_cells(row)
            if not ucells:
                continue
            row_text = " ".join(c.text for c in ucells)

            # --- Baris data utama IKU ---
            if penanda in ucells[0].text:
                for para in ucells[0].paragraphs:
                    for run in para.runs:
                        if penanda in run.text:
                            run.text = run.text.replace(penanda, kode_asli)
                            run.font.name = "Arial"
                target_pk_text = ("%s %s" % (iku["target_pk"], iku["satuan"])).strip()
                for pos, nilai in [(2, target_pk_text), (3, iku["alokasi_tw"]),
                                   (4, iku["real_tw"]), (5, iku["capaian_tw"]),
                                   (6, iku["capaian_pk"])]:
                    if len(ucells) > pos and ucells[pos].text.strip() == "":
                        set_cell_text(ucells[pos], nilai)

            # --- Realisasi Volume RO / Kendala / Solusi ---
            if "Kendala" in row_text and "Solusi" in row_text:
                for cell in ucells:
                    if "Kendala" in cell.text and "Solusi" in cell.text:
                        if iku["ro_narasi"]:
                            sisip_nilai_setelah_label(cell, "Realisasi Volume RO", iku["ro_narasi"])
                        for para in list(cell.paragraphs):
                            teks = para.text.strip()
                            if teks.startswith("Kendala") and iku["kendala"]:
                                para._p.addnext(buat_para_nilai(bersihkan_nilai(iku["kendala"])))
                            elif teks.startswith("Solusi") and iku["solusi"]:
                                para._p.addnext(buat_para_nilai(bersihkan_nilai(iku["solusi"])))

            # --- RTL & PIC ---
            if "Rencana Tindak Lanjut" in row_text and "PIC" in row_text:
                for cell in ucells:
                    if "Rencana Tindak Lanjut" in cell.text and "PIC" not in cell.text:
                        sisip_nilai_setelah_label(cell, "Rencana Tindak Lanjut",
                                                  bersihkan_nilai(iku["rtl"]))
                    if "PIC Tindak Lanjut" in cell.text:
                        sisip_nilai_setelah_label(cell, "Batas Waktu Tindak Lanjut",
                                                  bersihkan_nilai(iku["batas"]))
                        sisip_nilai_setelah_label(cell, "PIC Tindak Lanjut",
                                                  bersihkan_nilai(iku["pic"]))

            # --- Tautan bukti dukung realisasi IKU ---
            if "Tautan Bukti Dukung Realisasi IKU" in row_text and "Rencana" not in row_text:
                for cell in ucells:
                    if "Tautan Bukti Dukung Realisasi IKU" in cell.text:
                        sisip_nilai_setelah_label(cell, "Tautan Bukti Dukung Realisasi IKU",
                                                  bersihkan_nilai(iku["link_bukti"]), is_link=True)

            # --- Tautan bukti dukung RTL triwulan sebelumnya ---
            if "Tautan Bukti Dukung Rencana Tindak Lanjut Triwulan Sebelumnya" in row_text:
                for cell in ucells:
                    if "Tautan Bukti Dukung Rencana Tindak Lanjut" in cell.text:
                        sisip_nilai_setelah_label(cell, "Tautan Bukti Dukung Rencana Tindak Lanjut",
                                                  bersihkan_nilai(iku["link_tl_sblm"]), is_link=True)

    # ---- 4. FORMAT & SIMPAN ----
    set_arial_seluruh_doc(doc)
    doc.save(output_path)

    return {"iku_terisi": terisi, "iku_tanpa_tabel": tidak_ketemu}


# ============================================================
# MAIN
# ============================================================
def main():
    ap = argparse.ArgumentParser(description="Generate Notula Monitoring Kinerja")
    ap.add_argument("--excel",    help="File Excel kertas kerja yang di-upload")
    ap.add_argument("--sheet",    default="LK_Kabkot", help="Nama sheet Excel")
    ap.add_argument("--triwulan", help="Paksa triwulan: I / II / III / IV")
    ap.add_argument("--json",     help="File JSON data (alternatif Excel)")
    ap.add_argument("--meta",     help="File JSON metadata rapat (menimpa nilai Excel)")
    ap.add_argument("--template", help="File template word_FRA.docx")
    ap.add_argument("--output",   help="Path file hasil .docx")
    ap.add_argument("--inspect",  action="store_true",
                    help="Hanya baca sumber data lalu cetak JSON, tanpa membuat dokumen")
    ap.add_argument("--abaikan-cek", action="store_true",
                    help="Lewati pemeriksaan apakah template masih kosong")
    args = ap.parse_args()

    if not args.json and not args.excel:
        sys.stderr.write("ERROR: wajib pilih salah satu: --excel atau --json\n")
        sys.exit(2)
    if not args.inspect and (not args.template or not args.output):
        sys.stderr.write("ERROR: --template dan --output wajib diisi (kecuali mode --inspect)\n")
        sys.exit(2)

    for f in filter(None, [args.json, args.excel, args.meta, args.template]):
        if not os.path.exists(f):
            sys.stderr.write("ERROR: file tidak ditemukan: %s\n" % f)
            sys.exit(2)

    data = baca_json(args.json) if args.json else baca_excel(args.excel, args.sheet, args.triwulan)
    if args.meta:
        data = gabung_meta(data, args.meta)

    # ---- Mode inspect: kirim data mentah untuk preview di UI ----
    if args.inspect:
        print(json.dumps({
            "success": True,
            "mode": "inspect",
            "satker": data["satker"],
            "nilai_sakip": data["nilai_sakip"],
            "predikat": data["predikat"],
            "triwulan": data["triwulan"],
            "tahun": data.get("tahun", 2026),
            "sheet": data.get("_sheet", ""),
            "sheets_tersedia": data.get("_sheets_tersedia", []),
            "peringatan": data.get("_peringatan", []),
            "iku_list": data["iku_list"],
        }, ensure_ascii=False))
        return

    try:
        info = generate_notula(data, args.template, args.output, args.abaikan_cek)
    except Exception as e:
        sys.stderr.write("ERROR saat generate: %s\n" % e)
        sys.exit(3)

    print(json.dumps({
        "success": True,
        "mode": "generate",
        "output": args.output,
        "satker": data["satker"],
        "triwulan": data["triwulan"],
        "tahun": data.get("tahun", 2026),
        "sheet": data.get("_sheet", ""),
        "iku_total": len(data["iku_list"]),
        "iku_terisi": info["iku_terisi"],
        "iku_tanpa_tabel": info["iku_tanpa_tabel"],
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
